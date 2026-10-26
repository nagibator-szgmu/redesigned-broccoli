"""Вспомогательные функции форматирования и стилей для генерации docx документа."""
import re
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

NAVY_BLUE = RGBColor(26, 54, 93)     # #1A365D
DARK_RED = RGBColor(155, 44, 44)     # #9B2C2C
DARK_GRAY = RGBColor(74, 85, 104)    # #4A5568

EMOJI_PATTERN = re.compile(
    r"[\U00010000-\U0010ffff\u2600-\u27bf\u2300-\u23ff\u2b50\u2b55\u200d\ufe0f\u2700-\u27bf]"
)

def strip_emojis(text):
    """Полное удаление эмодзи и спецсимволов Юникода."""
    if not text:
        return ""
    cleaned = EMOJI_PATTERN.sub("", str(text))
    return re.sub(r" +", " ", cleaned).strip()

def set_cell_background(cell, hex_color):
    """Установить цвет фона ячейки таблицы."""
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def format_eff_dict(eff):
    """Форматирование словаря эффектов препарата."""
    labels = {
        "sbp": "САД", "dbp": "ДАД", "hr": "ЧСС", "rr": "ЧДД",
        "spo2": "SpO2", "temp": "Темп.", "gcs": "ШКГ", "pain": "Боль"
    }
    units = {
        "sbp": " мм рт.ст.", "dbp": " мм рт.ст.", "hr": " уд/мин",
        "rr": " в мин", "spo2": "%", "temp": " C", "gcs": " б", "pain": " б"
    }
    parts = []
    for k, val in (eff or {}).items():
        sign = "+" if val > 0 else ""
        name = labels.get(k, k)
        unit = units.get(k, "")
        parts.append(f"{name} {sign}{val}{unit}")
    return ", ".join(parts) if parts else "нет прямого витального сдвига"

def format_death_thresh(thresh):
    """Форматирование порогов гибели."""
    labels = {
        "sbp": "САД", "dbp": "ДАД", "hr": "ЧСС", "rr": "ЧДД",
        "spo2": "SpO2", "temp": "Темп.", "gcs": "ШКГ"
    }
    parts = []
    for k, v in (thresh or {}).items():
        name = labels.get(k, k)
        comp = ">=" if (k == "sbp" and v > 200) else "<="
        parts.append(f"{name} {comp} {v}")
    return ", ".join(parts) if parts else "Не задано"

def format_day_plan(d):
    """Форматирование дневника курации в стационаре."""
    if d.get("text"):
        return strip_emojis(d.get("text"))
    parts = []
    if d.get("day"):
        parts.append(f"День {d.get('day')}:")
    if d.get("morning"):
        parts.append(strip_emojis(str(d.get("morning"))))
    if d.get("plan"):
        parts.append(f"Назначения: {strip_emojis(d.get('plan'))}")
    if d.get("action"):
        parts.append(f"Действие: {strip_emojis(d.get('action'))}")
    return " ".join(parts)

def format_route_options(opts, correct_id):
    """Форматирование опций маршрутизации."""
    labels = []
    correct_label = correct_id
    for o in opts or []:
        if isinstance(o, dict):
            oid = o.get("id", "")
            lbl = strip_emojis(o.get("label", oid))
            labels.append(lbl)
            if oid == correct_id:
                correct_label = lbl
        else:
            txt = strip_emojis(str(o))
            labels.append(txt)
            if str(o) == correct_id:
                correct_label = txt
    return correct_label, ", ".join(labels)

def add_vitals_table(doc, c):
    """Добавить таблицу исходных витальных показателей."""
    v = c.get("vitals", {})
    t = doc.add_table(rows=2, cols=7)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["АД (мм рт.ст.)", "ЧСС (уд/м)", "ЧДД (в мин)", "SpO2", "t (C)", "ШКГ", "Боль (ВАШ)"]
    vals = [
        str(v.get("bp", "---")), str(v.get("hr", "---")), str(v.get("rr", "---")),
        f"{v.get('spo2', '---')}%", f"{v.get('temp', '---')} C",
        f"{c.get('initialGCS', '---')}/15", f"{c.get('initialPain', '0')}/10"
    ]
    for i, h in enumerate(headers):
        cell = t.cell(0, i); cell.text = h
        set_cell_background(cell, "2B6CB0")
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True; r.font.size = Pt(8.5); r.font.color.rgb = RGBColor(255, 255, 255)
    for i, val in enumerate(vals):
        cell = t.cell(1, i); cell.text = val
        set_cell_background(cell, "EDF2F7")
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.size = Pt(9.5); r.bold = True; r.font.color.rgb = NAVY_BLUE

def add_callout(doc, title, text, bg_hex="FFF5F5"):
    """Создать выделительную плашку без эмодзи."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_hex)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(3)
    run_t = p.add_run(f"{strip_emojis(title)}: "); run_t.bold = True; run_t.font.color.rgb = DARK_RED
    run_b = p.add_run(strip_emojis(text)); run_b.font.size = Pt(9); run_b.font.color.rgb = DARK_GRAY

def add_bullets(doc, items, prefix="- "):
    """Добавить ненумерованный список без эмодзи."""
    for it in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(f"{prefix}{strip_emojis(it)}")
        r.font.size = Pt(9)
