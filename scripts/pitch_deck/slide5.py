from pptx.util import Inches, Pt
from .config import (
    FONT_PRIMARY, C_DARK, C_MUTED, C_BLUE,
    C_BLUE_DARK, C_BLUE_BG, C_BLUE_BORDER,
    C_EMERALD, C_EMERALD_BG, C_EMERALD_BORDER,
    C_CARD_SLATE_BG, C_CARD_BORDER
)
from .helpers import (
    add_card, add_badge, set_tf_margins,
    clear_non_title_shapes, format_slide_title, add_bullets
)

def build_slide5(slide):
    clear_non_title_shapes(slide)
    format_slide_title(slide, "ЦА и СТЕЙКХОЛДЕРЫ. ПАРТНЕРЫ")

    w_col = Inches(4.98)
    left_c1 = Inches(1.40)
    left_c2 = Inches(6.62)

    # 1. Card 1: B2B / B2G Customer
    top_c1 = Inches(1.20)
    h_c1 = Inches(3.30)
    c1 = add_card(slide, left_c1, top_c1, w_col, h_c1)
    add_badge(slide, left_c1 + Inches(0.14), top_c1 + Inches(0.10),
              Inches(2.50), Inches(0.24), "01 / B2B & B2G: МЕДВУЗЫ И ЦЕНТРЫ",
              bg_color=C_BLUE_BG, border_color=C_BLUE_BORDER, text_color=C_BLUE, font_size=7.5)
    tf1 = c1.text_frame
    set_tf_margins(tf1, top=26, bottom=8, left=12, right=12)

    p = tf1.paragraphs[0]
    p.text = "120+ медицинских вузов, 110+ центров аккредитации, 400+ колледжей"
    p.font.name, p.font.size, p.font.bold = FONT_PRIMARY, Pt(9.5), True
    p.font.color.rgb = C_DARK
    p.space_after = Pt(4)

    add_bullets(tf1, [
        ("ЛПР (Принимают решение)", "Проректоры по учебной работе, руководители симуляционно-аттестационных центров, деканы факультетов."),
        ("Влияют на закупку", "Заведующие кафедрами терапии, анестезиологии, хирургии, преподаватели, актив СНО."),
        ("Ключевая ценность", "100% импортозамещение (замена Body Interact), выполнение ст. 37 ФЗ-323, подготовка к аккредитации (ОСКЭ), снижение затрат в 10 раз."),
        ("Модель монетизации", "Корпоративная годовая лицензия на вуз (от 1.2 до 3.5 млн ₽/год) в зависимости от контингента обучающихся.")
    ], font_size=8.5, space_after=3)

    # 2. Card 2: B2C Users
    top_c2 = Inches(4.62)
    h_c2 = Inches(2.18)
    c2 = add_card(slide, left_c1, top_c2, w_col, h_c2)
    add_badge(slide, left_c1 + Inches(0.14), top_c2 + Inches(0.10),
              Inches(2.50), Inches(0.24), "02 / B2C: СТУДЕНТЫ И ОРДИНАТОРЫ",
              bg_color=C_BLUE_BG, border_color=C_BLUE_BORDER, text_color=C_BLUE_DARK, font_size=7.5)
    tf2 = c2.text_frame
    set_tf_margins(tf2, top=26, bottom=8, left=12, right=12)

    p = tf2.paragraphs[0]
    p.text = "Студенты 3–6 курсов (лечебное дело, педиатрия), врачи-ординаторы"
    p.font.name, p.font.size, p.font.bold = FONT_PRIMARY, Pt(9.5), True
    p.font.color.rgb = C_DARK
    p.space_after = Pt(4)

    add_bullets(tf2, [
        ("Профиль и боль", "20–27 лет, РФ. Преодоление страха первого самостоятельного дежурства, дефицит практики у постели больного, подготовка к станциям ОСКЭ."),
        ("Модель монетизации", "Freemium (15 базовых кейсов бесплатно) + PRO-подписка (490–990 ₽/мес за полный банк 88 кейсов, дебрифинг и сертификаты).")
    ], font_size=8.5, space_after=3)

    # 3. Card 3: Partners
    h_c3 = Inches(2.65)
    c3 = add_card(slide, left_c2, top_c1, w_col, h_c3)
    add_badge(slide, left_c2 + Inches(0.14), top_c1 + Inches(0.10),
              Inches(2.40), Inches(0.24), "СТРАТЕГИЧЕСКИЕ ПАРТНЕРЫ",
              bg_color=C_EMERALD_BG, border_color=C_EMERALD_BORDER, text_color=C_EMERALD, font_size=7.5)
    tf3 = c3.text_frame
    set_tf_margins(tf3, top=26, bottom=8, left=12, right=12)

    add_bullets(tf3, [
        ("ФГБОУ ВО «СЗГМУ им. И.И. Мечникова» Минздрава РФ", "Клиническая база и академический соорганизатор. Врачебный аудит кейсов профессорами кафедры госпитальной терапии, симуляционный центр, площадка для пилотирования (70+ студентов)."),
        ("СПбГЭТУ «ЛЭТИ» (Акселератор «Технологический суверенитет»)", "Технологическая и бизнес-акселерация, трекинг продуктовых гипотез, инженерная IT-экспертиза, упаковка для рынка HealthNet НТИ, вывод на институты развития.")
    ], font_size=8.5, space_after=3)

    # 4. Card 4: Stakeholders
    top_c4 = Inches(3.97)
    h_c4 = Inches(2.83)
    c4 = add_card(slide, left_c2, top_c4, w_col, h_c4)
    add_badge(slide, left_c2 + Inches(0.14), top_c4 + Inches(0.10),
              Inches(2.40), Inches(0.24), "КЛЮЧЕВЫЕ СТЕЙКХОЛДЕРЫ",
              bg_color=C_CARD_SLATE_BG, border_color=C_CARD_BORDER, text_color=C_MUTED, font_size=7.5)
    tf4 = c4.text_frame
    set_tf_margins(tf4, top=26, bottom=8, left=12, right=12)

    add_bullets(tf4, [
        ("Минздрав России (Развивающее влияние)", "Закрепил обязательность соблюдения КР (ст. 37 ФЗ-323) и периодическую аккредитацию специалистов. Формирует нормативную рамку симуляционного обучения."),
        ("Минобрнауки РФ & Фонд содействия инновациям", "Финансовая и грантовая поддержка программ технологического суверенитета («Приоритет 2030», гранты «Студенческий стартап» 1 млн ₽)."),
        ("Руководство медвузов РФ (Ключевой закупщик)", "Формируют планы закупок симуляционного ПО. Заинтересованы в снижении издержек на зарубежные комплексы и повышении сдаваемости первичной аккредитации.")
    ], font_size=8.5, space_after=3)

    slide.notes_slide.notes_text_frame.text = (
        "СЛАЙД 5 | ЦА, СТЕЙКХОЛДЕРЫ И ПАРТНЕРЫ\n"
        "«Кто является нашим клиентом и партнером?\n"
        "В структуре продаж 85% выручки формирует сегмент B2B и B2G — это медицинские университеты Минздрава и Минобрнауки, а также симуляционные центры. "
        "Лица, принимающие решение, — это проректоры по учебной работе и руководители аккредитационных центров. Их главная боль — колоссальные затраты на зарубежные манекены, "
        "блокировка европейского софта и риск не сдать первичную аккредитацию выпускниками.\n"
        "Второй сегмент — B2C: сами студенты старших курсов и ординаторы, которые готовятся к ОСКЭ и дежурствам. Для них действует модель Freemium с подпиской на углубленные сценарии.\n"
        "Наши ключевые партнеры: во-первых, СЗГМУ им. Мечникова — это наш клинический дом, предоставивший кафедры госпитальной терапии, анестезиологии и симуляционный центр "
        "для верификации и пилотных испытаний на 70+ студентах. Во-вторых, СПбГЭТУ «ЛЭТИ» и программа «Технологический суверенитет» — это технологический и трекерский акселератор, "
        "помогающий упаковать проект для рынка HealthNet НТИ. Главные регуляторы — Минздрав и Минобрнауки — создали идеальные условия для внедрения отечественного симулятора.»"
    )
