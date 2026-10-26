from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from .config import (
    FONT_PRIMARY, CORNER_RADIUS, C_CARD_BG, C_CARD_BORDER,
    C_DARK, C_BODY, C_BLUE_BG, C_BLUE_BORDER, C_BLUE_DARK
)

def add_card(slide, left, top, width, height,
             bg_color=C_CARD_BG,
             border_color=C_CARD_BORDER,
             border_width=Pt(1),
             corner_radius=CORNER_RADIUS,
             shape_type=MSO_SHAPE.ROUNDED_RECTANGLE):
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    if corner_radius is not None and shape_type == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            shape.adjustments[0] = corner_radius
        except Exception:
            pass
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = border_width
    else:
        shape.line.fill.background()
    return shape

def add_badge(slide, left, top, width, height, text,
              bg_color=C_BLUE_BG, border_color=C_BLUE_BORDER,
              text_color=C_BLUE_DARK, font_size=8.0):
    card = add_card(slide, left, top, width, height,
                    bg_color=bg_color, border_color=border_color,
                    corner_radius=0.40)
    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_top = Pt(2)
    tf.margin_bottom = Pt(2)
    tf.margin_left = Pt(6)
    tf.margin_right = Pt(6)
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.CENTER
    p.font.name = FONT_PRIMARY
    p.font.size = Pt(font_size)
    p.font.bold = True
    p.font.color.rgb = text_color
    return card

def set_tf_margins(tf, top=10, bottom=10, left=12, right=12):
    tf.word_wrap = True
    tf.margin_top = Pt(top)
    tf.margin_bottom = Pt(bottom)
    tf.margin_left = Pt(left)
    tf.margin_right = Pt(right)

def remove_shape(slide, shape):
    slide.shapes._spTree.remove(shape._element)

def clear_non_title_shapes(slide):
    for sh in list(slide.shapes):
        if not sh.name.startswith('Заголовок'):
            remove_shape(slide, sh)

def format_slide_title(slide, title_text):
    for sh in slide.shapes:
        if sh.name.startswith('Заголовок') and sh.has_text_frame:
            sh.left = Inches(1.40)
            sh.top = Inches(0.28)
            sh.width = Inches(10.20)
            sh.height = Inches(0.82)
            tf = sh.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = title_text
            p.font.name = FONT_PRIMARY
            p.font.size = Pt(24)
            p.font.bold = True
            p.font.color.rgb = C_DARK
            break

def add_bullets(tf, items, font_size=9.5, space_after=4):
    for title, text in items:
        p = tf.add_paragraph()
        p.space_after = Pt(space_after)
        r1 = p.add_run()
        r1.text = "• " + title + ": "
        r1.font.name = FONT_PRIMARY
        r1.font.bold = True
        r1.font.size = Pt(font_size)
        r1.font.color.rgb = C_DARK
        r2 = p.add_run()
        r2.text = text
        r2.font.name = FONT_PRIMARY
        r2.font.size = Pt(font_size)
        r2.font.color.rgb = C_BODY
