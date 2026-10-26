from pptx.util import Inches, Pt
from .config import (
    FONT_PRIMARY, C_DARK, C_MUTED, C_BLUE, C_BLUE_DARK,
    C_BLUE_BG, C_BLUE_BORDER, C_EMERALD, C_EMERALD_BG, C_EMERALD_BORDER
)
from .helpers import (
    add_card, add_badge, set_tf_margins,
    clear_non_title_shapes, format_slide_title, add_bullets
)

def build_slide3(slide):
    clear_non_title_shapes(slide)
    format_slide_title(slide, "РЕШЕНИЕ и ТЕХНОЛОГИЧЕСКОЕ ЯДРО ПРОЕКТА")

    left_base = Inches(1.40)
    w_left = Inches(4.98)
    left_right = Inches(6.58)
    w_right = Inches(5.02)
    top_base = Inches(1.20)

    # 1. Left Card: Product Showcase with Real UI Preview
    c_left = add_card(slide, left_base, top_base, w_left, Inches(4.15))
    add_badge(slide, left_base + Inches(0.14), top_base + Inches(0.12),
              Inches(3.20), Inches(0.26), "РАБОЧАЯ СТАНЦИЯ ВРАЧА • MEDSIM",
              bg_color=C_BLUE_BG, border_color=C_BLUE_BORDER, text_color=C_BLUE_DARK)

    # Embed real high-res screenshot
    screenshot_path = "reports/screenshots/desktop/18_game_icu_workspace.png"
    slide.shapes.add_picture(
        screenshot_path,
        left_base + Inches(0.14), top_base + Inches(0.48),
        width=w_left - Inches(0.28), height=Inches(2.70)
    )

    # Bottom caption inside left card
    tf_l = c_left.text_frame
    set_tf_margins(tf_l, top=238, bottom=6, left=12, right=12)
    p = tf_l.paragraphs[0]
    p.text = "4 клинических отделения: ОРИТ • Приёмное • Поликлиника • Стационар"
    p.font.name, p.font.size, p.font.bold = FONT_PRIMARY, Pt(9), True
    p.font.color.rgb = C_DARK
    p.space_after = Pt(2)

    p2 = tf_l.add_paragraph()
    p2.text = "Сквозная клиническая траектория: от 12-минутной реанимации до суточного обхода и амбулаторной маршрутизации."
    p2.font.name, p2.font.size = FONT_PRIMARY, Pt(8)
    p2.font.color.rgb = C_MUTED

    # 2. Right Top Card: Mathematical Deterioration Engine
    c_r1 = add_card(slide, left_right, top_base, w_right, Inches(2.00))
    add_badge(slide, left_right + Inches(0.14), top_base + Inches(0.10),
              Inches(2.60), Inches(0.24), "01 / МАТЕМАТИЧЕСКАЯ ФИЗИОЛОГИЯ",
              bg_color=C_BLUE_BG, border_color=C_BLUE_BORDER, text_color=C_BLUE_DARK, font_size=8)
    tf_r1 = c_r1.text_frame
    set_tf_margins(tf_r1, top=28, bottom=6, left=12, right=12)

    p = tf_r1.paragraphs[0]
    p.text = "30 секунд — такт пересчета динамики ЖП"
    p.font.name, p.font.size, p.font.bold = FONT_PRIMARY, Pt(11), True
    p.font.color.rgb = C_BLUE_DARK
    p.space_after = Pt(3)

    add_bullets(tf_r1, [
        ("Динамическая физиология", "Математическое моделирование 7 ЖП (АД, ЧСС, SpO2, ЧДД, темп, ШКГ, SOFA) в реальном времени."),
        ("Фармакодинамика 60 препаратов", "Учет времени начала (onset), пика эффекта, кумуляции и клиренса. Ошибочный выбор ведет к шоку и смерти."),
        ("«Слепая» лаборатория", "Студент сам выбирает из 38 методов и оценивает биохимию (15 параметров) и ОАМ без подсказок системы.")
    ], font_size=8.5, space_after=2)

    # 3. Right Bottom Card: Critical Technology & 11-point Debriefing
    c_r2 = add_card(slide, left_right, top_base + Inches(2.12), w_right, Inches(2.03))
    add_badge(slide, left_right + Inches(0.14), top_base + Inches(2.20),
              Inches(2.60), Inches(0.24), "02 / КРИТИЧЕСКАЯ ТЕХНОЛОГИЯ РФ",
              bg_color=C_EMERALD_BG, border_color=C_EMERALD_BORDER, text_color=C_EMERALD, font_size=8)
    tf_r2 = c_r2.text_frame
    set_tf_margins(tf_r2, top=26, bottom=6, left=12, right=12)

    p = tf_r2.paragraphs[0]
    p.text = "100% КР — 88 протоколов cr.minzdrav.gov.ru"
    p.font.name, p.font.size, p.font.bold = FONT_PRIMARY, Pt(11), True
    p.font.color.rgb = C_EMERALD
    p.space_after = Pt(3)

    add_bullets(tf_r2, [
        ("11-точечный дебрифинг", "Автоматический аудит диагностической траектории, пропущенных находок, полипрагмазии и ятрогенных ошибок."),
        ("Критическая технология РФ", "«Создание доверенного отечественного ПО и моделирование биосистем» (HealthNet НТИ)."),
        ("Суверенный стек по 152-ФЗ", "100% браузерный запуск, zero-dependency от зарубежных серверов и облаков.")
    ], font_size=8.5, space_after=2)

    # 4. Bottom Canonical Formula Banner
    banner = add_card(slide, left_base, Inches(5.45), Inches(10.20), Inches(1.35),
                      bg_color=C_EMERALD_BG, border_color=C_EMERALD_BORDER)
    tf_b = banner.text_frame
    set_tf_margins(tf_b, top=8, bottom=8, left=14, right=14)

    pb0 = tf_b.paragraphs[0]
    pb0.text = "КАНОНИЧЕСКАЯ ФОРМУЛИРОВКА РЕШЕНИЯ (ШАБЛОН АКСЕЛЕРАТОРА):"
    pb0.font.name, pb0.font.size, pb0.font.bold = FONT_PRIMARY, Pt(9), True
    pb0.font.color.rgb = C_EMERALD
    pb0.space_after = Pt(2)

    pb1 = tf_b.add_paragraph()
    r = pb1.add_run()
    r.text = (
        "«Используя отечественный математический детериорационный движок динамической физиологии и алгоритмическую валидацию "
        "по КР Минздрава РФ, MedSim устраняет дефицит практики, повышая клиническую точность решений студентов на 35% "
        "и сокращая затраты медвузов на симуляционное обучение более чем в 10 раз.»"
    )
    r.font.name, r.font.size, r.font.color.rgb = FONT_PRIMARY, Pt(9.5), C_DARK

    slide.notes_slide.notes_text_frame.text = (
        "СЛАЙД 3 (3:00 - 6:00) | РЕШЕНИЕ И ТЕХНОЛОГИЧЕСКОЕ ЯДРО\n"
        "«Что представляет собой MedSim? Это не статичный квиз, а полноценная цифровая рабочая станция врача.\n"
        "Студент курирует пациентов в 4 отделениях: от экстренного ОРИТ с живым монитором ЭКГ и давления до многодневного терапевтического стационара.\n"
        "В основе лежит наше технологическое ядро — динамический детериорационный движок. Каждые 30 секунд система пересчитывает физиологические показатели: "
        "АД, ЧСС, сатурацию, шкалу Глазго с учётом введенных препаратов. Если студент медлит или назначает противопоказанные препараты — начинается острая декомпенсация.\n"
        "Наша уникальная авторская механика — «слепая» лаборатория: студент сам заказывает биохимию и мочу и сам выносит вердикт по отклонениям, система не дает спойлеров.\n"
        "А в финале алгоритм формирует 11-точечный дебрифинг — детальный аудит решений по стандартам Минздрава РФ. Это полностью суверенная разработка, "
        "относящаяся к критической технологии моделирования биологических систем.»"
    )
