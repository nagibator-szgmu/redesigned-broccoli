from pptx.util import Inches, Pt
from .config import FONT_PRIMARY, C_DARK, C_BLUE, C_BLUE_DARK
from .helpers import format_slide_title

def build_slide6(slide):
    """
    SLIDE 6: Value Proposition.
    Elevator pitch & accelerator canonical value formula.
    """
    format_slide_title(slide, "ЦЕННОСТНОЕ ПРЕДЛОЖЕНИЕ")
    if len(slide.shapes) > 1:
        shape6 = slide.shapes[1]
        shape6.left = Inches(1.40)
        shape6.top = Inches(1.30)
        shape6.width = Inches(10.20)
        shape6.height = Inches(5.20)
        tf6 = shape6.text_frame
        tf6.clear()

        slide6_blocks = [
            ("ФОРМУЛА ЦЕННОСТНОГО ПРЕДЛОЖЕНИЯ (ШАБЛОН АКСЕЛЕРАТОРА):", True, 12, C_BLUE),
            ("«Наш интерактивный симулятор-тренажер MedSim помогает медицинским вузам, симуляционным центрам и студентам-медикам, которые хотят эффективно освоить практические навыки экстренной помощи и подготовиться к аккредитации (ОСКЭ),", False, 10, C_DARK),
            ("тем, что погружает их в реалистичное клиническое дежурство с динамическим ухудшением состояния пациента и пошаговым разбором ошибок по стандартам Минздрава РФ,", False, 10, C_DARK),
            ("и снижает стоимость подготовки одного специалиста более чем в 10 раз по сравнению с аппаратными манекенами,", False, 10, C_DARK),
            ("(в отличие от зарубежных симуляторов Body Interact и SimX, которые стоят миллионы рублей, заблокированы в РФ и противоречат российским клиническим рекомендациям).»", True, 10, C_BLUE_DARK),
            ("", False, 8, C_DARK),
            ("КЛЮЧЕВОЙ ELEVATOR PITCH (ОДНА ФРАЗА):", True, 12, C_BLUE),
            ("«MedSim — это безопасный тренажер клинического мышления, превращающий студента в уверенного врача до его первой встречи с реальным тяжелым пациентом.»", True, 11, C_DARK)
        ]

        for i, (text, bold, size, color) in enumerate(slide6_blocks):
            p = tf6.paragraphs[0] if i == 0 else tf6.add_paragraph()
            p.text = text
            p.font.name = FONT_PRIMARY
            p.font.bold = bold
            p.font.size = Pt(size)
            p.font.color.rgb = color
            p.space_after = Pt(4)

def build_slide7(slide):
    """
    SLIDE 7: Competitors & Tech Sovereignty Table.
    """
    format_slide_title(slide, "КОНКУРЕНТЫ и АНАЛОГИ. РЫНОК")
    for sh in slide.shapes:
        if sh.has_table:
            sh.left = Inches(1.40)
            sh.top = Inches(1.50)
            sh.width = Inches(10.20)
            table7 = sh.table
            table7.cell(0, 1).text = "MedSim\n(СЗГМУ / ЛЭТИ, РФ)"
            table7.cell(0, 2).text = "Body Interact\n(Португалия/США)"
            table7.cell(0, 3).text = "Clinical Sense\n(США)"
            table7.cell(0, 4).text = "Манекены Laerdal\n(Норвегия)"

            row1_data = [
                "Соответствие КР Минздрава РФ и ФЗ-323",
                "100% (88 кейсов по cr.minzdrav.gov.ru)",
                "0% (только гайдлайны США/ЕС)",
                "0% (гайдлайны США, FDA)",
                "Частично (зависит от очного тьютора)"
            ]
            for c_i, val in enumerate(row1_data):
                table7.cell(1, c_i).text = val

            row2_data = [
                "Технологический суверенитет и стоимость",
                "Российский стек, 0 санкционных рисков, веб 24/7",
                "Лицензии заблокированы в РФ, валютный риск",
                "Заблокирован, только базовые квизы",
                "Импорт от 25 млн ₽/комплекс, очереди в центрах"
            ]
            for c_i, val in enumerate(row2_data):
                table7.cell(2, c_i).text = val
            break
