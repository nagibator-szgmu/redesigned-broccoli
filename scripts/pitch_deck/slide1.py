from pptx.util import Inches, Pt
from pptx.enum.text import MSO_ANCHOR
from .config import (
    FONT_PRIMARY, C_DARK, C_MUTED, C_BLUE,
    C_BLUE_DARK, C_BLUE_BG, C_BLUE_BORDER, C_CARD_BG, C_CARD_BORDER
)
from .helpers import add_card, add_badge, set_tf_margins, clear_non_title_shapes

def build_slide1(slide):
    """
    SLIDE 1: Title & Accelerator Context.
    Strictly preserves user manual edits (Mikhail Goloshchapov, Victoria Kichova, Irina Rasmagina).
    Positioned cleanly to the right of the cloud illustration (left >= 5.65").
    """
    clear_non_title_shapes(slide)

    left_content = Inches(5.65)
    w_content = Inches(6.05)

    # Category chip positioned cleanly in clear space
    add_badge(
        slide, left_content, Inches(2.22), Inches(4.80), Inches(0.28),
        "HEALTHNET НТИ  •  СИМУЛЯЦИОННОЕ ОБУЧЕНИЕ  •  ИМПОРТОЗАМЕЩЕНИЕ",
        bg_color=C_BLUE_BG, border_color=C_BLUE_BORDER, text_color=C_BLUE_DARK, font_size=7.0
    )

    title_shape = slide.shapes[0]
    title_shape.left = left_content
    title_shape.top = Inches(2.58)
    title_shape.width = w_content
    title_shape.height = Inches(1.55)
    title_shape.text_frame.clear()
    title_shape.text_frame.vertical_anchor = MSO_ANCHOR.TOP

    p0 = title_shape.text_frame.paragraphs[0]
    p0.text = "MedSim"
    p0.font.name = FONT_PRIMARY
    p0.font.bold = True
    p0.font.size = Pt(36)
    p0.font.color.rgb = C_DARK
    p0.space_after = Pt(2)

    p1 = title_shape.text_frame.add_paragraph()
    p1.text = "Интерактивный симулятор клинического мышления на основе математического моделирования физиологии и стандартов Минздрава РФ"
    p1.font.name = FONT_PRIMARY
    p1.font.size = Pt(11.5)
    p1.font.bold = True
    p1.font.color.rgb = C_BLUE_DARK
    p1.space_after = Pt(3)

    p2 = title_shape.text_frame.add_paragraph()
    p2.text = "Акселерационная программа СПбГЭТУ «ЛЭТИ» «Технологический суверенитет» (Осень 2026)"
    p2.font.name = FONT_PRIMARY
    p2.font.size = Pt(9.5)
    p2.font.bold = True
    p2.font.color.rgb = C_MUTED

    # Team & Institution Card Container
    card_top = Inches(4.30)
    card_h = Inches(2.55)
    card = add_card(slide, left_content, card_top, w_content, card_h,
                    bg_color=C_CARD_BG, border_color=C_CARD_BORDER)

    add_badge(slide, left_content + Inches(0.14), card_top + Inches(0.08),
              Inches(2.20), Inches(0.22), "ПАСПОРТ СТАРТАП-ПРОЕКТА",
              bg_color=C_BLUE_BG, border_color=C_BLUE_BORDER, text_color=C_BLUE_DARK, font_size=7.5)

    tf = card.text_frame
    tf.vertical_anchor = MSO_ANCHOR.TOP
    set_tf_margins(tf, top=24, bottom=6, left=12, right=12)

    meta_lines = [
        ("Организация-разработчик: ", "ФГБОУ ВО «СЗГМУ им. И.И. Мечникова» Минздрава России"),
        ("Лидер стартап-проекта: ", "Голощапов Михаил (студент 3 курса лечебного факультета СЗГМУ им. Мечникова)"),
        ("Трекер проекта: ", "Кичова Виктория (СПбГЭТУ «ЛЭТИ»)"),
        ("Наставник проекта: ", "Расмагина Ирина Алексеевна (как Вас записать лучше всего?"),
        ("Куратор направления: ", "Направление HealthNet / ИИ в медицине, СПбГЭТУ «ЛЭТИ»")
    ]

    for i, (label, val) in enumerate(meta_lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(3.5)
        r1 = p.add_run()
        r1.text = label
        r1.font.name = FONT_PRIMARY
        r1.font.bold = True
        r1.font.size = Pt(8.5)
        r1.font.color.rgb = C_MUTED
        r2 = p.add_run()
        r2.text = val
        r2.font.name = FONT_PRIMARY
        r2.font.bold = False
        r2.font.size = Pt(8.5)
        r2.font.color.rgb = C_DARK

    slide.notes_slide.notes_text_frame.text = (
        "СЛАЙД 1 (0:00 - 1:00) | ТИТУЛЬНЫЙ\n"
        "«Добрый день, уважаемые эксперты и члены комиссии! Меня зовут Михаил Голощапов, я студент 3 курса лечебного "
        "факультета СЗГМУ им. Мечникова и лидер стартап-проекта MedSim.\n"
        "Мы представляем проект в рамках акселерационной программы СПбГЭТУ «ЛЭТИ» «Технологический суверенитет» от Северо-Западного государственного "
        "медицинского университета им. Мечникова.\n"
        "Сегодня каждый студент-медик сталкивается с парадоксом: мы учим терапию и реанимацию по статичным тестам. "
        "Но когда молодой врач впервые заходит в палату ОРИТ или приёмного покоя, перед ним живой тяжелый пациент, состояние которого "
        "ухудшается каждую минуту. MedSim — это отечественный интерактивный симулятор клинического мышления, "
        "который учит будущего врача действовать в условиях цейтнота строго по Клиническим рекомендациям Минздрава РФ без риска для жизни пациентов.»"
    )
