from pptx.util import Inches, Pt
from .config import (
    FONT_PRIMARY, C_DARK, C_MUTED, C_BLUE, C_BLUE_BG, C_BLUE_BORDER,
    C_EMERALD, C_EMERALD_BG, C_EMERALD_BORDER, C_CARD_SLATE_BG, C_CARD_BORDER
)
from .helpers import (
    add_card, add_badge, set_tf_margins,
    clear_non_title_shapes, format_slide_title, add_bullets
)

def build_slide4(slide):
    clear_non_title_shapes(slide)
    format_slide_title(slide, "УРОВЕНЬ ГОТОВНОСТИ ТЕХНОЛОГИИ (TRL 5–6)")

    left_base = Inches(1.40)
    w_total = Inches(10.20)

    # 1. Top Status Banner
    banner = add_card(slide, left_base, Inches(1.15), w_total, Inches(0.68),
                      bg_color=C_EMERALD_BG, border_color=C_EMERALD_BORDER)
    tf_b = banner.text_frame
    set_tf_margins(tf_b, top=8, bottom=8, left=14, right=14)

    p0 = tf_b.paragraphs[0]
    p0.text = "ТЕКУЩИЙ УРОВЕНЬ ГОТОВНОСТИ ТЕХНОЛОГИИ: TRL 5 (С ЭЛЕМЕНТАМИ TRL 6)"
    p0.font.name, p0.font.size, p0.font.bold = FONT_PRIMARY, Pt(11), True
    p0.font.color.rgb = C_EMERALD
    p0.space_after = Pt(2)

    p1 = tf_b.add_paragraph()
    r = p1.add_run()
    r.text = (
        "Полнофункциональный лабораторный образец испытан в релевантной среде (веб-комплекс 4 отделений) • "
        "Опытная эксплуатация прототипа на кафедре госпитальной терапии и в СНО СЗГМУ им. Мечникова"
    )
    r.font.name, r.font.size, r.font.color.rgb = FONT_PRIMARY, Pt(9), C_DARK

    # 2. Row of 4 Metric Cards
    metrics = [
        ("88", "Клинических сценариев", "38 ОРИТ • 27 приёмное • 16 поликлиника • 7 стационар", C_BLUE),
        ("60 / 38", "Лекарств и диагностик", "60 препаратов с фармакодинамикой • 38 методов обследования", C_BLUE),
        ("100% КР", "Стандарты Минздрава РФ", "Полное соответствие рубрикатору cr.minzdrav.gov.ru", C_EMERALD),
        ("70+ чел.", "Пилотная апробация", "Выборка студентов СЗГМУ им. Мечникова (+35% точности)", C_EMERALD)
    ]
    card_w = Inches(2.40)
    gap = Inches(0.20)
    top_m = Inches(1.92)
    h_m = Inches(1.02)

    for i, (stat, title, sub, color) in enumerate(metrics):
        left_m = left_base + i * (card_w + gap)
        c = add_card(slide, left_m, top_m, card_w, h_m)
        tf = c.text_frame
        set_tf_margins(tf, top=6, bottom=6, left=10, right=10)

        p = tf.paragraphs[0]
        p.text = stat
        p.font.name, p.font.size, p.font.bold = FONT_PRIMARY, Pt(18), True
        p.font.color.rgb = color

        p = tf.add_paragraph()
        p.text = title
        p.font.name, p.font.size, p.font.bold = FONT_PRIMARY, Pt(9), True
        p.font.color.rgb = C_DARK

        p = tf.add_paragraph()
        p.text = sub
        p.font.name, p.font.size = FONT_PRIMARY, Pt(7.5)
        p.font.color.rgb = C_MUTED

    # 3. Two Evidence Cards (TRL 5 vs TRL 6)
    top_e = Inches(3.04)
    h_e = Inches(2.32)
    w_e = Inches(4.98)

    # Card A: TRL 5 Artifacts
    ca = add_card(slide, left_base, top_e, w_e, h_e)
    add_badge(slide, left_base + Inches(0.14), top_e + Inches(0.10),
              Inches(3.20), Inches(0.24), "АРТЕФАКТЫ TRL 5 (ЛАБОРАТОРНЫЙ ОБРАЗЕЦ)",
              bg_color=C_BLUE_BG, border_color=C_BLUE_BORDER, text_color=C_BLUE, font_size=8)
    tfa = ca.text_frame
    set_tf_margins(tfa, top=28, bottom=8, left=12, right=12)

    add_bullets(tfa, [
        ("Действующий веб-симулятор MedSim", "Полнофункциональная кроссплатформенная рабочая станция врача, развернутая на тестовом стенде."),
        ("Математическое ядро детериорации", "Непрерывный пересчет витальных показателей каждые 30 сек с учетом фармакодинамики и полиорганной недостаточности."),
        ("Инженерный сьют верификации", "100% покрытие авто-валидаторами, 20+ e2e тестов Playwright, 0 ошибок формул, строгий клинический интерфейс.")
    ], font_size=8.5, space_after=3)

    # Card B: TRL 6 Artifacts
    left_cb = left_base + w_e + Inches(0.24)
    cb = add_card(slide, left_cb, top_e, w_e, h_e)
    add_badge(slide, left_cb + Inches(0.14), top_e + Inches(0.10),
              Inches(3.20), Inches(0.24), "АРТЕФАКТЫ TRL 6 (ОПЫТНАЯ ЭКСПЛУАТАЦИЯ)",
              bg_color=C_EMERALD_BG, border_color=C_EMERALD_BORDER, text_color=C_EMERALD, font_size=8)
    tfb = cb.text_frame
    set_tf_margins(tfb, top=28, bottom=8, left=12, right=12)

    add_bullets(tfb, [
        ("Клиническая база испытаний", "Кафедра госпитальной терапии и СНО СЗГМУ им. И.И. Мечникова Минздрава России."),
        ("Протоколы пилотных апробаций", "Успешные испытания на выборке из 70+ студентов старших курсов с подтвержденным ростом точности диагнозов на 35%."),
        ("Официальный реестр верификации", "88 клинических паспортов сценариев сопоставлены с утвержденными КР Минздрава РФ (cr.minzdrav.gov.ru).")
    ], font_size=8.5, space_after=3)

    # 4. Bottom Roadmap / TRL Milestones
    road = add_card(slide, left_base, Inches(5.45), w_total, Inches(1.35),
                    bg_color=C_CARD_SLATE_BG, border_color=C_CARD_BORDER)
    tf_r = road.text_frame
    set_tf_margins(tf_r, top=8, bottom=8, left=14, right=14)

    p0 = tf_r.paragraphs[0]
    p0.text = "СООТВЕТСТВИЕ КРИТЕРИЯМ АКСЕЛЕРАТОРА И ГОРИЗОНТ РАЗВИТИЯ TRL:"
    p0.font.name, p0.font.size, p0.font.bold = FONT_PRIMARY, Pt(9), True
    p0.font.color.rgb = C_MUTED
    p0.space_after = Pt(2)

    add_bullets(tf_r, [
        ("TRL 1–4 [Завершено]", "Концепция, математические модели динамики гемодинамики, лабораторный веб-прототип движка."),
        ("TRL 5–6 [ТЕКУЩИЙ СТАТУС]", "Полнофункциональная рабочая станция в релевантной среде, 88 клинических сценариев, пилотирование на 70+ студентах."),
        ("TRL 7 [Цель акселератора 2026–2027]", "Опытно-промышленное внедрение в симуляционные центры 3 медвузов РФ, регистрация программы ЭВМ в Роспатенте.")
    ], font_size=8.5, space_after=2)

    slide.notes_slide.notes_text_frame.text = (
        "СЛАЙД 4 (6:00 - 7:30) | УРОВЕНЬ ГОТОВНОСТИ (TRL 5-6)\n"
        "«На какой стадии проект сегодня? Мы находимся на уровне TRL 5 с переходом в TRL 6.\n"
        "Это не презентация идеи, а готовый работающий программный комплекс. В системе развернуто 88 клинических кейсов "
        "по 7 ключевым специальностям: от инфаркта миокарда и ТЭЛА до острой кишечной непроходимости и цирроза печени. "
        "Все кейсы прошли медицинский аудит и сопоставлены с рубрикатором Минздрава.\n"
        "Наши артефакты: действующий веб-тренажер, который мы готовы продемонстрировать в браузере прямо сейчас; "
        "полный сьют автоматических тестов; реестр клинического ревью и акты пилотных апробаций в СНО кафедры госпитальной терапии СЗГМУ им. Мечникова "
        "на выборке более 70 студентов, показавшие рост клинической точности на 35%.»"
    )
