import json
import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_cell_background(cell, fill_hex):
    """Sets the background color of a table cell."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    """Sets inner padding for a table cell in dxa (1 pt = 20 dxa)."""
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    """Sets subtle light borders on all table cells."""
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def add_callout(doc, text, title=""):
    """Adds a stylish callout box with a colored left accent border."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F0FDF4") # soft emerald tint
    set_cell_margins(cell, top=140, bottom=140, left=220, right=200)
    
    # Left border only (accent)
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="059669"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    if title:
        run_title = p.add_run(f"{title}\n")
        run_title.bold = True
        run_title.font.name = "Arial"
        run_title.font.size = Pt(10)
        run_title.font.color.rgb = RGBColor(4, 120, 87)
    
    run_text = p.add_run(text)
    run_text.italic = True
    run_text.font.name = "Arial"
    run_text.font.size = Pt(10.5)
    run_text.font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def main():
    json_path = "scripts/gastro_cases_data.json"
    with open(json_path, "r", encoding="utf-8") as f:
        cases = json.load(f)

    doc = Document()

    # Page setup: Standard margins (2 cm)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles
    styles = doc.styles
    normal_style = styles['Normal']
    normal_style.font.name = 'Arial'
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = RGBColor(30, 41, 59) # Slate 800

    # Colors
    PRIMARY = RGBColor(15, 118, 110)   # Teal 700
    DARK_BLUE = RGBColor(30, 58, 138)  # Blue 900
    MUTED = RGBColor(100, 116, 139)    # Slate 500

    # ── TITLE / HEADER ────────────────────────────────────────────────────────
    p_badge = doc.add_paragraph()
    p_badge.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_badge = p_badge.add_run("MEDSIM® CLINICAL SIMULATION PLATFORM • РГА / АКР / МИНЗДРАВ РФ 2024")
    r_badge.font.size = Pt(8.5)
    r_badge.font.bold = True
    r_badge.font.color.rgb = PRIMARY
    p_badge.paragraph_format.space_after = Pt(4)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("Амбулаторная гастроэнтерология\n10 клинических кейсов поликлиники")
    r_title.font.size = Pt(20)
    r_title.font.bold = True
    r_title.font.color.rgb = DARK_BLUE
    p_title.paragraph_format.space_after = Pt(8)

    p_desc = doc.add_paragraph()
    p_desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_desc = p_desc.add_run(
        "Полное методическое руководство по амбулаторно-поликлиническому приёму (Outpatient Department).\n"
        "Жалобы, дифференциальная диагностика, лабораторно-инструментальные критерии, базисная фармакотерапия и маршрутизация по Клиническим рекомендациям Минздрава России."
    )
    r_desc.font.size = Pt(10.5)
    r_desc.font.italic = True
    r_desc.font.color.rgb = MUTED
    p_desc.paragraph_format.space_after = Pt(18)

    # ── SUMMARY TABLE ─────────────────────────────────────────────────────────
    p_tbl_title = doc.add_paragraph()
    r_tbl_title = p_tbl_title.add_run("Сводный реестр 10 гастроэнтерологических задач поликлиники")
    r_tbl_title.font.size = Pt(12)
    r_tbl_title.font.bold = True
    r_tbl_title.font.color.rgb = PRIMARY
    p_tbl_title.paragraph_format.space_after = Pt(6)

    sum_table = doc.add_table(rows=len(cases) + 1, cols=6)
    sum_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(sum_table)

    headers = ["№", "ID", "Пациент", "Клинический диагноз", "МКБ-10", "Маршрут"]
    col_widths = [Inches(0.4), Inches(1.1), Inches(1.2), Inches(2.2), Inches(0.9), Inches(1.1)]

    hdr_row = sum_table.rows[0]
    for i, title in enumerate(headers):
        cell = hdr_row.cells[i]
        cell.width = col_widths[i]
        set_cell_background(cell, "0F766E")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(title)
        r.font.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)

    for idx, c in enumerate(cases):
        row = sum_table.rows[idx + 1]
        bg = "F8FAFC" if idx % 2 == 1 else "FFFFFF"
        route_text = "Амбулаторно" if c["correctRoute"] == "treat_outpatient" else "К хирургу (ЛХЭ)" if c["correctRoute"] == "refer_specialist" else "Госпитализация"
        
        cells_data = [
            str(idx + 1),
            c["id"],
            f"{c['name']}\n({c['age']} л., {c['gender']})",
            c["diagnosis"],
            c["icd"].split(" ")[0],
            route_text
        ]

        for i, val in enumerate(cells_data):
            cell = row.cells[i]
            cell.width = col_widths[i]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in [0, 1, 4, 5] else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(val)
            r.font.size = Pt(8.5)
            if i in [1, 5]:
                r.font.bold = True
                if i == 5 and "хирург" in val:
                    r.font.color.rgb = RGBColor(180, 83, 9)

    doc.add_page_break()

    # ── 10 DETAILED CASES ─────────────────────────────────────────────────────
    for idx, c in enumerate(cases):
        # Case Header
        p_c_title = doc.add_paragraph()
        r_c_num = p_c_title.add_run(f"КЕЙС {idx + 1}: ")
        r_c_num.font.size = Pt(14)
        r_c_num.font.bold = True
        r_c_num.font.color.rgb = PRIMARY

        r_c_name = p_c_title.add_run(f"{c['name']} ({c['age']} лет, {c['gender'] == 'М' and 'мужчина' or 'женщина'})")
        r_c_name.font.size = Pt(14)
        r_c_name.font.bold = True
        r_c_name.font.color.rgb = DARK_BLUE
        p_c_title.paragraph_format.space_after = Pt(4)

        # Meta table
        meta_table = doc.add_table(rows=4, cols=2)
        meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(meta_table)

        meta_rows = [
            ("Идентификатор (ID) / МКБ-10:", f"{c['id']}  •  {c['icd']}"),
            ("Клинический диагноз:", c['diagnosis']),
            ("Клинические рекомендации:", f"{c['sourceReference']['name']} ({c['sourceReference']['year']} г.)"),
            ("Допустимые формулировки:", ", ".join([f"«{v}»" for v in c['diagnosisVariants']]) if c['diagnosisVariants'] else "—")
        ]

        for r_idx, (k, v) in enumerate(meta_rows):
            cell_k = meta_table.cell(r_idx, 0)
            cell_v = meta_table.cell(r_idx, 1)
            cell_k.width = Inches(2.2)
            cell_v.width = Inches(4.6)
            set_cell_background(cell_k, "F1F5F9")
            set_cell_background(cell_v, "FFFFFF")
            set_cell_margins(cell_k, top=60, bottom=60, left=100, right=100)
            set_cell_margins(cell_v, top=60, bottom=60, left=100, right=100)
            
            pk = cell_k.paragraphs[0]
            pk.paragraph_format.space_after = Pt(0)
            rk = pk.add_run(k)
            rk.font.bold = True
            rk.font.size = Pt(9)
            rk.font.color.rgb = DARK_BLUE

            pv = cell_v.paragraphs[0]
            pv.paragraph_format.space_after = Pt(0)
            rv = pv.add_run(v)
            rv.font.size = Pt(9)
            if r_idx == 1:
                rv.font.bold = True

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

        # 1. Complaints
        p_sec1 = doc.add_paragraph()
        r_s1 = p_sec1.add_run("1. Жалобы пациента при обращении")
        r_s1.font.bold = True
        r_s1.font.size = Pt(11)
        r_s1.font.color.rgb = PRIMARY
        p_sec1.paragraph_format.space_after = Pt(2)

        add_callout(doc, f"«{c['complaint']}»", "Первичный анамнестический фокус:")

        # 2. Anamnesis
        p_sec2 = doc.add_paragraph()
        r_s2 = p_sec2.add_run("2. Анамнестические данные")
        r_s2.font.bold = True
        r_s2.font.size = Pt(11)
        r_s2.font.color.rgb = PRIMARY
        p_sec2.paragraph_format.space_after = Pt(2)

        p_morbi = doc.add_paragraph()
        r_m_tag = p_morbi.add_run("• Анамнез заболевания (Anamnesis morbi): ")
        r_m_tag.font.bold = True
        r_m_tag.font.size = Pt(9.5)
        r_m_text = p_morbi.add_run(c['historyOfIllness'])
        r_m_text.font.size = Pt(9.5)
        p_morbi.paragraph_format.space_after = Pt(3)

        p_vitae = doc.add_paragraph()
        r_v_tag = p_vitae.add_run("• Анамнез жизни и факторы риска (Anamnesis vitae): ")
        r_v_tag.font.bold = True
        r_v_tag.font.size = Pt(9.5)
        r_v_text = p_vitae.add_run(c['lifeHistory'] or "Аллергологический анамнез не отягощен. Вредных привычек нет.")
        r_v_text.font.size = Pt(9.5)
        p_vitae.paragraph_format.space_after = Pt(6)

        # 3. Status & Vitals
        p_sec3 = doc.add_paragraph()
        r_s3 = p_sec3.add_run("3. Физикальный статус и показатели жизнедеятельности")
        r_s3.font.bold = True
        r_s3.font.size = Pt(11)
        r_s3.font.color.rgb = PRIMARY
        p_sec3.paragraph_format.space_after = Pt(2)

        # Vitals mini-table
        vit_table = doc.add_table(rows=2, cols=6)
        vit_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(vit_table)
        
        v_headers = ["АД", "ЧСС", "ЧДД", "T тела", "SpO₂", "Боль / ШКГ"]
        v_vals = [
            f"{c['vitals'].get('bp', '120/80')} мм",
            f"{c['vitals'].get('hr', 72)} уд/мин",
            f"{c['vitals'].get('rr', 16)} /мин",
            f"{c['vitals'].get('temp', 36.6)} °C",
            f"{c['vitals'].get('spo2', 99)}%",
            f"{c['initialPain']}/10 • {c['initialGCS']} б."
        ]
        
        for vi in range(6):
            c_hdr = vit_table.cell(0, vi)
            c_val = vit_table.cell(1, vi)
            c_hdr.width = Inches(1.1)
            c_val.width = Inches(1.1)
            set_cell_background(c_hdr, "0F766E")
            set_cell_background(c_val, "F8FAFC")
            set_cell_margins(c_hdr, top=60, bottom=60, left=80, right=80)
            set_cell_margins(c_val, top=60, bottom=60, left=80, right=80)
            
            p_vh = c_hdr.paragraphs[0]
            p_vh.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_vh.paragraph_format.space_after = Pt(0)
            r_vh = p_vh.add_run(v_headers[vi])
            r_vh.font.bold = True
            r_vh.font.size = Pt(8.5)
            r_vh.font.color.rgb = RGBColor(255, 255, 255)

            p_vv = c_val.paragraphs[0]
            p_vv.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_vv.paragraph_format.space_after = Pt(0)
            r_vv = p_vv.add_run(v_vals[vi])
            r_vv.font.bold = True
            r_vv.font.size = Pt(8.5)
            r_vv.font.color.rgb = DARK_BLUE

        p_exam = doc.add_paragraph()
        p_exam.paragraph_format.space_before = Pt(4)
        r_ex_tag = p_exam.add_run("Объективный осмотр: ")
        r_ex_tag.font.bold = True
        r_ex_tag.font.size = Pt(9.5)
        r_ex_text = p_exam.add_run(c['exam'] or "Общее состояние удовлетворительное. Живот мягкий, пальпация без особенностей. Перитонеальных знаков нет.")
        r_ex_text.font.size = Pt(9.5)
        p_exam.paragraph_format.space_after = Pt(6)

        # 4. Diagnostics Table
        p_sec4 = doc.add_paragraph()
        r_s4 = p_sec4.add_run("4. Диагностический план и результаты обследований")
        r_s4.font.bold = True
        r_s4.font.size = Pt(11)
        r_s4.font.color.rgb = PRIMARY
        p_sec4.paragraph_format.space_after = Pt(2)

        d_table = doc.add_table(rows=len(c['testResults']) + 1, cols=3)
        d_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(d_table)

        d_hdr = d_table.rows[0]
        for di, dh_text in enumerate(["Исследование", "Статус", "Клинический результат исследования"]):
            cell = d_hdr.cells[di]
            cell.width = [Inches(1.8), Inches(1.1), Inches(3.9)][di]
            set_cell_background(cell, "0F766E")
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if di == 1 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(dh_text)
            r.font.bold = True
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor(255, 255, 255)

        for ti, t in enumerate(c['testResults']):
            row = d_table.rows[ti + 1]
            bg = "FFFFFF" if ti % 2 == 0 else "F8FAFC"
            
            c0 = row.cells[0]
            c0.width = Inches(1.8)
            set_cell_background(c0, bg)
            set_cell_margins(c0, top=60, bottom=60, left=100, right=100)
            p0 = c0.paragraphs[0]
            p0.paragraph_format.space_after = Pt(0)
            r0 = p0.add_run(t['name'])
            r0.font.bold = True
            r0.font.size = Pt(8.5)

            c1 = row.cells[1]
            c1.width = Inches(1.1)
            set_cell_background(c1, bg)
            set_cell_margins(c1, top=60, bottom=60, left=100, right=100)
            p1 = c1.paragraphs[0]
            p1.paragraph_format.space_after = Pt(0)
            p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r1 = p1.add_run("Обязательный" if t['isRequired'] else "Дополнительно")
            r1.font.bold = t['isRequired']
            r1.font.size = Pt(8)
            r1.font.color.rgb = RGBColor(4, 120, 87) if t['isRequired'] else MUTED

            c2 = row.cells[2]
            c2.width = Inches(3.9)
            set_cell_background(c2, bg)
            set_cell_margins(c2, top=60, bottom=60, left=100, right=100)
            p2 = c2.paragraphs[0]
            p2.paragraph_format.space_after = Pt(0)
            r2 = p2.add_run(t['result'])
            r2.font.size = Pt(8.5)
            if "🔴" in t['result'] or "язва" in t['result'].lower() or "конкремент" in t['result'].lower():
                r2.font.bold = True

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

        # 5. Treatment
        p_sec5 = doc.add_paragraph()
        r_s5 = p_sec5.add_run("5. Лечебная тактика и фармакотерапия")
        r_s5.font.bold = True
        r_s5.font.size = Pt(11)
        r_s5.font.color.rgb = PRIMARY
        p_sec5.paragraph_format.space_after = Pt(2)

        if not c['needTreat']:
            p_nt = doc.add_paragraph()
            r_nt = p_nt.add_run("• Рекомендованная терапия: Немедикаментозная коррекция (диета, снижения массы тела, нормализация физической активности, исключение гепатотоксинов).")
            r_nt.font.size = Pt(9.5)
            r_nt.font.italic = True
        else:
            for nt in c['needTreat']:
                p_nt = doc.add_paragraph()
                p_nt.paragraph_format.space_after = Pt(2)
                r_nt_icon = p_nt.add_run("✓ ")
                r_nt_icon.font.bold = True
                r_nt_icon.font.color.rgb = RGBColor(16, 185, 129)
                r_nt_name = p_nt.add_run(f"{nt['name']}: ")
                r_nt_name.font.bold = True
                r_nt_name.font.size = Pt(9.5)
                r_nt_desc = p_nt.add_run(nt['note'])
                r_nt_desc.font.size = Pt(9.5)

        for wt in c['wrongTreat']:
            p_wt = doc.add_paragraph()
            p_wt.paragraph_format.space_after = Pt(2)
            r_wt_icon = p_wt.add_run("✗ Опасно! ")
            r_wt_icon.font.bold = True
            r_wt_icon.font.color.rgb = RGBColor(220, 38, 38)
            r_wt_name = p_wt.add_run(f"{wt['name']}: ")
            r_wt_name.font.bold = True
            r_wt_name.font.size = Pt(9.5)
            r_wt_desc = p_wt.add_run(wt['reason'])
            r_wt_desc.font.size = Pt(9.5)

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

        # 6. Routing
        p_sec6 = doc.add_paragraph()
        r_s6 = p_sec6.add_run("6. Маршрутизация пациента")
        r_s6.font.bold = True
        r_s6.font.size = Pt(11)
        r_s6.font.color.rgb = PRIMARY
        p_sec6.paragraph_format.space_after = Pt(2)

        corr_opt = next((ro for ro in c['routeOptions'] if ro['isCorrect']), None)
        p_route_main = doc.add_paragraph()
        r_rm_tag = p_route_main.add_run("Правильное решение: ")
        r_rm_tag.font.bold = True
        r_rm_tag.font.size = Pt(9.5)
        r_rm_tag.font.color.rgb = RGBColor(4, 120, 87)
        r_rm_text = p_route_main.add_run(corr_opt['label'] if corr_opt else c['correctRoute'])
        r_rm_text.font.bold = True
        r_rm_text.font.size = Pt(9.5)

        if corr_opt and corr_opt['reason']:
            p_route_exp = doc.add_paragraph()
            p_route_exp.paragraph_format.space_after = Pt(4)
            r_re = p_route_exp.add_run(f"Обоснование: {corr_opt['reason']}")
            r_re.font.italic = True
            r_re.font.size = Pt(9)
            r_re.font.color.rgb = MUTED

        # 7. Debrief
        p_sec7 = doc.add_paragraph()
        r_s7 = p_sec7.add_run("7. Академический дебрифинг и ключевые стандарты Минздрава РФ")
        r_s7.font.bold = True
        r_s7.font.size = Pt(11)
        r_s7.font.color.rgb = PRIMARY
        p_sec7.paragraph_format.space_after = Pt(2)

        if c['debrief']:
            p_deb = doc.add_paragraph()
            p_deb.paragraph_format.space_after = Pt(4)
            r_deb_tag = p_deb.add_run("Клинический разбор: ")
            r_deb_tag.font.bold = True
            r_deb_tag.font.size = Pt(9.5)
            r_deb_text = p_deb.add_run(c['debrief'])
            r_deb_text.font.size = Pt(9.5)

        if c['checklistItems']:
            p_chk_hdr = doc.add_paragraph()
            p_chk_hdr.paragraph_format.space_after = Pt(2)
            r_ch = p_chk_hdr.add_run("Чек-лист обязательных действий врача:")
            r_ch.font.bold = True
            r_ch.font.size = Pt(9.5)
            for ch_item in c['checklistItems']:
                p_chi = doc.add_paragraph()
                p_chi.paragraph_format.space_after = Pt(1)
                r_box = p_chi.add_run("☑ ")
                r_box.font.color.rgb = PRIMARY
                r_txt = p_chi.add_run(ch_item)
                r_txt.font.size = Pt(9)

        if idx < len(cases) - 1:
            doc.add_page_break()

    # Save to both target locations
    out1 = "c:/Users/мишка/Desktop/redesigned-broccoli-main/redesigned-broccoli/docs/Клинические_кейсы_поликлиники_Гастроэнтерология_MedSim.docx"
    out2 = "c:/Users/мишка/Desktop/Клинические_кейсы_поликлиники_Гастроэнтерология_MedSim.docx"
    
    doc.save(out1)
    doc.save(out2)
    print(f"Successfully generated DOCX files at:\n1. {out1}\n2. {out2}")

if __name__ == "__main__":
    main()
