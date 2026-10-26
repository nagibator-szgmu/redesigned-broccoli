"""Основной генератор документа Word (.docx) со всеми 79 клиническими кейсами MedSim без эмодзи."""
import json
import os
import pathlib
from docx import Document
from docx.shared import Inches, Pt

from docx_builder import (
    NAVY_BLUE, DARK_GRAY, format_eff_dict, format_death_thresh,
    format_day_plan, format_route_options, strip_emojis,
    add_vitals_table, add_callout, add_bullets
)

BUNDLE_PATH = pathlib.Path(__file__).parent / "cases_bundle.json"
DESKTOP_PATH = pathlib.Path(os.environ.get("USERPROFILE", "C:/Users/мишка")) / "Desktop"
OUTPUT_FILE = DESKTOP_PATH / "Клинические_кейсы_MedSim.docx"

CATEGORIES = [
    ("cardiac", "I. КАРДИОЛОГИЯ"),
    ("respiratory", "II. ПУЛЬМОНОЛОГИЯ"),
    ("gastro", "III. ГАСТРОЭНТЕРОЛОГИЯ"),
    ("neuro", "IV. НЕВРОЛОГИЯ"),
    ("infectious", "V. ИНФЕКЦИОННЫЕ БОЛЕЗНИ"),
    ("endocrine", "VI. ЭНДОКРИНОЛОГИЯ"),
    ("toxicology", "VII. ТОКСИКОЛОГИЯ"),
]

DEPT_NAMES = {
    "icu": "ОРИТ (Реанимация)", "admission": "Приёмное отделение",
    "outpatient": "Поликлиника", "stationary": "Стационар"
}
SEV_NAMES = {"critical": "Критическая", "moderate": "Средняя", "mild": "Лёгкая"}

def render_case(doc, c, diags, treats, treat_fx, adv_fx, adv_reasons, notes):
    """Отрисовка одного клинического кейса в документ без эмодзи."""
    name_clean = strip_emojis(c['name'])
    h2 = doc.add_heading(f"Кейс {c['id']}: {name_clean} ({c.get('age', '—')} лет, {c.get('gender', '—')})", level=2)
    h2.paragraph_format.space_before = Pt(8); h2.paragraph_format.space_after = Pt(2)
    for r in h2.runs: r.font.color.rgb = NAVY_BLUE

    dept = DEPT_NAMES.get(c.get("department"), c.get("department", "—"))
    sev = SEV_NAMES.get(c.get("severity"), c.get("severity", "—"))
    tlim = f"{c.get('timeLimit')} мин." if c.get("timeLimit") else "не ограничен"
    ref = c.get("sourceReference", {})
    ref_txt = f"{strip_emojis(ref.get('name', 'КР РФ'))} ({ref.get('year', 2024)})" if ref else "КР Минздрава РФ"

    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(4)
    r1 = p_meta.add_run(f"Отделение: {dept}  |  Тяжесть: {sev}  |  Лимит времени: {tlim}\nКлинические рекомендации: {ref_txt}")
    r1.font.size = Pt(8.5); r1.font.italic = True; r1.font.color.rgb = DARK_GRAY

    p_clin = doc.add_paragraph()
    p_clin.paragraph_format.space_after = Pt(4)
    p_clin.add_run("Жалобы: ").bold = True; p_clin.add_run(f"{strip_emojis(c.get('complaint', '—'))}\n")
    hist = c.get("anamnesis") or c.get("historyOfIllness") or c.get("shortHistory") or "—"
    p_clin.add_run("Анамнез заболевания: ").bold = True; p_clin.add_run(f"{strip_emojis(hist)}\n")
    if c.get("lifeHistory"):
        p_clin.add_run("Анамнез жизни: ").bold = True; p_clin.add_run(f"{strip_emojis(c.get('lifeHistory'))}\n")
    p_clin.add_run("Объективный осмотр: ").bold = True; p_clin.add_run(f"{strip_emojis(c.get('exam', '—'))}")

    add_vitals_table(doc, c)

    if c.get("deterioration") or c.get("deathThresholds"):
        det_txt = (
            f"Темп декомпенсации (в мин): {format_eff_dict(c.get('deterioration'))}. "
            f"Пороги гибели: {format_death_thresh(c.get('deathThresholds'))}."
        )
        add_callout(doc, "ДИНАМИКА УХУДШЕНИЯ БЕЗ ТЕРАПИИ", det_txt, bg_hex="FFF5F5")

    p_diag = doc.add_paragraph()
    p_diag.paragraph_format.space_before = Pt(4); p_diag.paragraph_format.space_after = Pt(2)
    p_diag.add_run("Основной диагноз: ").bold = True
    r_diag = p_diag.add_run(strip_emojis(c.get("diagnosis", "—"))); r_diag.bold = True; r_diag.font.color.rgb = NAVY_BLUE
    if c.get("diagnosisVariants"):
        p_v = doc.add_paragraph(); p_v.paragraph_format.space_after = Pt(3)
        p_v.add_run("Допустимые варианты формулировки: ").bold = True
        p_v.add_run("; ".join([strip_emojis(v) for v in c.get("diagnosisVariants", [])]))

    doc.add_paragraph("Необходимые диагностические исследования:").runs[0].bold = True
    res_map = c.get("testResults", {})
    add_bullets(doc, [f"{strip_emojis(diags.get(did, did))}: {strip_emojis(res_map.get(did, 'Выполнено'))}" for did in c.get("needDiag", [])])

    doc.add_paragraph("Лечебные назначения и фармакодинамика:").runs[0].bold = True
    treat_lines = []
    for tid in c.get("needTreat", []):
        tname = strip_emojis(treats.get(tid, tid)); fx = treat_fx.get(tid, {})
        eff_str = format_eff_dict(fx.get("eff"))
        del_str = f"начало через {fx.get('delay')} с" if fx.get("delay") else "быстрое"
        cont_str = "непрерывно" if fx.get("continuous") else "болюсно"
        desc = strip_emojis(fx.get("desc", "")); note = strip_emojis(notes.get(tid, ""))
        treat_lines.append(f"{tname} — Сдвиг витальных функций: [{eff_str}], {del_str}, {cont_str}. {desc}. ({note})")
    add_bullets(doc, treat_lines)

    if c.get("checklistItems"):
        doc.add_paragraph("Чек-лист обязательных действий врача:").runs[0].bold = True
        add_bullets(doc, [strip_emojis(it) for it in c.get("checklistItems")], prefix="- ")

    if c.get("wrongTreat"):
        wrong_lines = []
        for wid in c.get("wrongTreat", []):
            wname = strip_emojis(treats.get(wid, wid)); afx = adv_fx.get(wid, {})
            afx_str = format_eff_dict(afx) if afx else "угроза осложнений"
            reason = strip_emojis(adv_reasons.get(wid, "Ухудшает клинический прогноз."))
            wrong_lines.append(f"{wname} — Опасный сдвиг: [{afx_str}]. Обоснование: {reason}")
        add_callout(doc, "ПРОТИВОПОКАЗАНИЯ И ОШИБКИ", "\n- " + "\n- ".join(wrong_lines))

    if c.get("department") == "outpatient" and c.get("correctRoute"):
        c_lbl, opts_str = format_route_options(c.get("routeOptions"), c.get("correctRoute"))
        add_callout(doc, "МАРШРУТИЗАЦИЯ ПАЦИЕНТА", f"Верный маршрут: {c_lbl}. Доступные опции: {opts_str}", bg_hex="EBF8FF")
    elif c.get("department") == "stationary" and c.get("dayByDayPlan"):
        doc.add_paragraph(f"Курация в стационаре (макс. {c.get('maxDays', '—')} дней):").runs[0].bold = True
        add_bullets(doc, [format_day_plan(d) for d in c.get("dayByDayPlan", [])])
        if c.get("dischargeCriteria"):
            add_bullets(doc, [strip_emojis(cr) for cr in c.get("dischargeCriteria")], prefix="- Критерий выписки: ")

    doc.add_page_break()

def main():
    node_export_cmd = pathlib.Path(__file__).parent / "export_cases_data.mjs"
    os.system(f"node {node_export_cmd}")
    with open(BUNDLE_PATH, "r", encoding="utf-8") as f:
        data = json.load(f)

    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.6); s.bottom_margin = Inches(0.6)
        s.left_margin = Inches(0.6); s.right_margin = Inches(0.6)

    cases = data["cases"]
    diags = {d["id"]: d["name"] for d in data["diagnostics"]}
    treats = {t["id"]: t["name"] for t in data["treatments"]}
    treat_fx = data["treatFx"]; adv_fx = data["adverseFx"]
    adv_reasons = data["adverseReasons"]; notes = data["treatNotes"]

    for cat_id, cat_title in CATEGORIES:
        cat_cases = [c for c in cases if c.get("category") == cat_id]
        if not cat_cases:
            continue
        h1 = doc.add_heading(f"{cat_title} ({len(cat_cases)} кейсов)", level=1)
        h1.paragraph_format.space_before = Pt(12); h1.paragraph_format.space_after = Pt(8)
        for r in h1.runs:
            r.font.color.rgb = NAVY_BLUE; r.bold = True

        for c in cat_cases:
            render_case(doc, c, diags, treats, treat_fx, adv_fx, adv_reasons, notes)

    doc.save(str(OUTPUT_FILE))
    if BUNDLE_PATH.exists():
        BUNDLE_PATH.unlink()
    print(f"OK: Документ успешно создан без эмодзи: {OUTPUT_FILE} ({len(cases)} кейсов)")

if __name__ == "__main__":
    main()
