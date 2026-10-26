import pptx
from pitch_deck import (
    build_slide1,
    build_slide2,
    build_slide3,
    build_slide4,
    build_slide5,
    build_slide6,
    build_slide7
)

def update_presentation():
    output_path = '/Users/arsenykozlov/Desktop/проекты/redesigned-broccoli-main/docs/MedSim_Pitch_LETI_Autumn_2026.pptx'
    prs = pptx.Presentation(output_path)

    print("Building Slide 1 (Title & Team)...")
    build_slide1(prs.slides[0])

    print("Building Slide 2 (Relevance & Problem)...")
    build_slide2(prs.slides[1])

    print("Building Slide 3 (Solution & Tech Core)...")
    build_slide3(prs.slides[2])

    print("Building Slide 4 (TRL 5-6 Readiness & Artifacts)...")
    build_slide4(prs.slides[3])

    print("Building Slide 5 (Target Audience, Stakeholders & Partners)...")
    build_slide5(prs.slides[4])

    print("Building Slide 6 (Value Proposition)...")
    build_slide6(prs.slides[5])

    print("Building Slide 7 (Competitors & Sovereignty Table)...")
    build_slide7(prs.slides[6])

    prs.save(output_path)
    print(f"Successfully generated updated pitch deck at: {output_path}")

if __name__ == '__main__':
    update_presentation()
